
import sqlite3
from pathlib import Path
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# --------------------------------------------------
# 1. Find the SQLite database and output Excel file
# --------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent
DATABASE_FILE = BASE_DIR / "database.db"
EXCEL_FILE = BASE_DIR / "blood_bank_database.xlsx"


def safe_sheet_name(name, used_names):
    """Create a valid, unique Excel worksheet name."""
    invalid_chars = '[]:*?/\\'
    cleaned = "".join(
        "_" if char in invalid_chars else char
        for char in name
    ).strip("'")

    cleaned = cleaned[:31] or "Table"
    candidate = cleaned
    number = 1

    while candidate.casefold() in used_names:
        suffix = f"_{number}"
        candidate = cleaned[:31 - len(suffix)] + suffix
        number += 1

    used_names.add(candidate.casefold())
    return candidate


def export_database():
    # --------------------------------------------------
    # 2. Check whether the database file exists
    # --------------------------------------------------

    if not DATABASE_FILE.is_file():
        print("ERROR: database.db was not found.")
        print("Expected location:", DATABASE_FILE)
        print("Place this script beside your database.db file.")
        return

    connection = None
    workbook = Workbook()
    workbook.remove(workbook.active)

    try:
        # Open the existing database in read-only mode
        connection = sqlite3.connect(
            DATABASE_FILE.as_uri() + "?mode=ro",
            uri=True
        )
        connection.row_factory = sqlite3.Row

        # --------------------------------------------------
        # 3. Get all user-created tables
        # --------------------------------------------------

        tables = connection.execute("""
            SELECT name
            FROM sqlite_master
            WHERE type = 'table'
              AND name NOT LIKE 'sqlite_%'
            ORDER BY name
        """).fetchall()

        if not tables:
            print("No database tables were found.")
            return

        used_sheet_names = set()

        # --------------------------------------------------
        # 4. Export each database table to its own worksheet
        # --------------------------------------------------

        for table_record in tables:
            table_name = table_record["name"]

            # Escape quotes in the table name safely
            quoted_name = '"' + table_name.replace('"', '""') + '"'

            cursor = connection.execute(
                f"SELECT * FROM {quoted_name}"
            )

            column_names = [
                description[0]
                for description in cursor.description
            ]

            worksheet_name = safe_sheet_name(
                table_name, used_sheet_names
            )
            worksheet = workbook.create_sheet(worksheet_name)

            # Write column headings
            worksheet.append(column_names)

            # Write every row from the database
            row_count = 0

            for record in cursor:
                worksheet.append([
                    record[column]
                    for column in column_names
                ])
                row_count += 1

            # Keep the heading row visible while scrolling
            worksheet.freeze_panes = "A2"

            # Format column headings
            for cell in worksheet[1]:
                cell.font = Font(
                    bold=True,
                    color="FFFFFF"
                )
                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=True
                )

            worksheet.row_dimensions[1].height = 24

            # Create an Excel table with filter dropdowns
            last_column = get_column_letter(len(column_names))
            last_row = max(1, worksheet.max_row)
            table_range = f"A1:{last_column}{last_row}"

            excel_table = Table(
                displayName=f"Table{len(used_sheet_names)}",
                ref=table_range
            )

            table_style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False
            )

            excel_table.tableStyleInfo = table_style
            worksheet.add_table(excel_table)

            # Adjust column widths to fit their contents
            for column_index, column_name in enumerate(
                column_names, start=1
            ):
                max_length = len(str(column_name))

                for row in worksheet.iter_rows(
                    min_col=column_index,
                    max_col=column_index,
                    min_row=2
                ):
                    value = row[0].value

                    if value is not None:
                        max_length = max(
                            max_length, len(str(value))
                        )

                worksheet.column_dimensions[
                    get_column_letter(column_index)
                ].width = min(max_length + 2, 40)

            print(
                f"Exported: {table_name} "
                f"({row_count} records)"
            )

        # --------------------------------------------------
        # 5. Save the Excel workbook
        # --------------------------------------------------

        workbook.save(EXCEL_FILE)

        print("\nEXPORT COMPLETED SUCCESSFULLY!")
        print("Excel file:", EXCEL_FILE)
        print("Worksheets created:", len(tables))

    except sqlite3.Error as error:
        print("DATABASE ERROR:", error)

    except PermissionError:
        print(
            "PERMISSION ERROR: Close the Excel file if it is "
            "already open, then run this program again."
        )

    except Exception as error:
        print("EXPORT ERROR:", error)

    finally:
        if connection is not None:
            connection.close()


if __name__ == "__main__":
    export_database()
