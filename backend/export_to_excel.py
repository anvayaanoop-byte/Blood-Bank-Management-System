
import os
import sqlite3
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================
# Blood Bank Management System
# Database to Excel Export
# ============================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATABASE = os.path.join(BASE_DIR, "database.db")

# Preferred output filename
OUTPUT = os.path.join(
    BASE_DIR, "blood_bank_database_export.xlsx"
)


def get_output_path():
    """
    Use the preferred filename if it is available.
    If locked, try a timestamped filename.
    """
    candidates = [
        OUTPUT,
        os.path.join(
            BASE_DIR,
            "blood_bank_database_"
            + datetime.now().strftime("%Y%m%d_%H%M%S")
            + ".xlsx"
        ),
        os.path.join(
            os.path.expanduser("~"),
            "Documents",
            "blood_bank_database_export.xlsx"
        ),
    ]

    for path in candidates:
        try:
            # Test whether the directory is writable.
            folder = os.path.dirname(path)
            if not os.path.isdir(folder):
                continue

            test_path = path + ".write_test"
            with open(test_path, "w", encoding="utf-8") as file:
                file.write("test")
            os.remove(test_path)

            # If the existing file is locked, skip it.
            if os.path.exists(path):
                try:
                    with open(path, "ab"):
                        pass
                except PermissionError:
                    continue

            return path

        except (PermissionError, OSError):
            continue

    raise PermissionError(
        "Cannot write the Excel file. Close Excel, "
        "check OneDrive sync, and try again."
    )


def style_header(sheet):
    """Format worksheet headers."""
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="C00000"
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


def autofit_columns(sheet):
    """Set readable column widths."""
    for column_cells in sheet.columns:
        max_length = 0
        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length, len(str(cell.value))
                )

        sheet.column_dimensions[column_letter].width = min(
            max_length + 2, 35
        )


def export_database():
    """Export each SQLite table into its own Excel worksheet."""

    if not os.path.exists(DATABASE):
        print("ERROR: database.db was not found.")
        print("Expected location:", DATABASE)
        return

    conn = None

    try:
        conn = sqlite3.connect(DATABASE)
        cursor = conn.cursor()

        tables = cursor.execute("""
            SELECT name
            FROM sqlite_master
            WHERE type = 'table'
              AND name NOT LIKE 'sqlite_%'
            ORDER BY name
        """).fetchall()

        if not tables:
            print("No database tables found.")
            return

        workbook = Workbook()
        summary = workbook.active
        summary.title = "Database Summary"

        summary.append(["Blood Bank Management System"])
        summary.append(["Database Table", "Record Count"])

        for cell in summary[2]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="C00000"
            )

        exported_count = 0

        for (table_name,) in tables:
            # Table names are retrieved from SQLite metadata.
            quoted_name = '"' + table_name.replace('"', '""') + '"'
            cursor.execute(f"SELECT * FROM {quoted_name}")

            rows = cursor.fetchall()
            columns = [
                description[0]
                for description in cursor.description
            ]

            # Excel worksheet names cannot exceed 31 characters.
            sheet_name = table_name[:31]

            # Avoid duplicate worksheet names.
            original_name = sheet_name
            suffix = 1
            while sheet_name in workbook.sheetnames:
                sheet_name = f"{original_name[:27]}_{suffix}"
                suffix += 1

            sheet = workbook.create_sheet(title=sheet_name)
            sheet.append(columns)

            for row in rows:
                sheet.append(list(row))

            style_header(sheet)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            autofit_columns(sheet)

            summary.append([table_name, len(rows)])
            exported_count += 1

        summary.freeze_panes = "A3"
        summary.column_dimensions["A"].width = 30
        summary.column_dimensions["B"].width = 18
        summary.auto_filter.ref = f"A2:B{summary.max_row}"

        output_path = get_output_path()
        workbook.save(output_path)

        print("\n====================================")
        print("EXCEL EXPORT COMPLETED SUCCESSFULLY")
        print("====================================")
        print("Database:", DATABASE)
        print("Excel file:", output_path)
        print("Tables exported:", exported_count)
        print("Total worksheets:", len(workbook.sheetnames))
        print("\nWorksheets:")

        for sheet_name in workbook.sheetnames:
            print("-", sheet_name)

    except sqlite3.Error as error:
        print("Database error:", error)

    except PermissionError as error:
        print("Permission error:", error)
        print("Close Excel and check folder permissions.")

    except Exception as error:
        print("Export failed:", error)

    finally:
        if conn is not None:
            conn.close()


if __name__ == "__main__":
    export_database()
